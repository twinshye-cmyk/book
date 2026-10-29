import io
import os
import sqlite3
from datetime import datetime, timezone
from functools import wraps
from pathlib import Path
from zoneinfo import ZoneInfo

from flask import Flask, abort, flash, redirect, render_template, request, send_file, session, url_for
from flask_wtf.csrf import CSRFProtect
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from werkzeug.security import check_password_hash, generate_password_hash
from markupsafe import escape

app = Flask(__name__)
secret = os.environ.get('SECRET_KEY')
if not secret or len(secret) < 32:
    raise RuntimeError('Set SECRET_KEY to a random value of at least 32 characters')
app.config.update(SECRET_KEY=secret, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.getenv('COOKIE_SECURE','1')=='1', MAX_CONTENT_LENGTH=128*1024)
CSRFProtect(app)
DATA_DIR = Path(os.getenv('DATA_DIR', './data'))
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB = DATA_DIR / 'submissions.db'
ADMIN_PASSWORD_HASH = os.getenv('ADMIN_PASSWORD_HASH','')
if not ADMIN_PASSWORD_HASH:
    raise RuntimeError('Set ADMIN_PASSWORD_HASH using generate_password_hash')

@app.after_request
def security_headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    response.headers['Content-Security-Policy'] = "default-src 'self'; style-src 'self'; script-src 'self'; base-uri 'self'; form-action 'self'; frame-ancestors 'none'"
    response.headers['Cache-Control'] = 'no-store'
    return response

def connect():
    con = sqlite3.connect(DB, timeout=15)
    con.row_factory = sqlite3.Row
    return con

def init_db():
    with connect() as con:
        con.execute('''CREATE TABLE IF NOT EXISTS submissions (id INTEGER PRIMARY KEY AUTOINCREMENT, student_id TEXT NOT NULL, student_name TEXT NOT NULL, book_title TEXT NOT NULL, author TEXT NOT NULL, review TEXT NOT NULL, created_at TEXT NOT NULL)''')
init_db()

def admin_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get('admin'):
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapped

def get_rows():
    q = request.args.get('q','').strip()[:100]
    with connect() as con:
        if q:
            like = '%' + q.replace('\\','\\\\').replace('%','\\%').replace('_','\\_') + '%'
            rows = con.execute("SELECT * FROM submissions WHERE student_id LIKE ? ESCAPE '\\' OR student_name LIKE ? ESCAPE '\\' OR book_title LIKE ? ESCAPE '\\' OR author LIKE ? ESCAPE '\\' ORDER BY id DESC", (like,)*4).fetchall()
        else:
            rows = con.execute('SELECT * FROM submissions ORDER BY id DESC').fetchall()
    return rows, q

@app.get('/')
def index():
    return render_template('student.html')

@app.post('/submit')
def submit():
    fields = {k:request.form.get(k,'').strip() for k in ('student_id','student_name','book_title','author','review')}
    limits = {'student_id':30,'student_name':60,'book_title':200,'author':120,'review':10000}
    if any(not fields[k] or len(fields[k]) > limits[k] for k in fields) or len(fields['review']) < 200:
        flash('모든 항목을 입력하고 독후감을 200자 이상 10,000자 이하로 작성해 주세요.', 'error')
        return redirect(url_for('index'))
    stamp = datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M:%S')
    with connect() as con:
        con.execute('INSERT INTO submissions(student_id,student_name,book_title,author,review,created_at) VALUES (?,?,?,?,?,?)', (*fields.values(),stamp))
    return redirect(url_for('submitted'))

@app.get('/submitted')
def submitted():
    return render_template('submitted.html')

@app.route('/teacher/login', methods=['GET','POST'])
def login():
    if request.method == 'POST':
        pw = request.form.get('password','')
        if check_password_hash(ADMIN_PASSWORD_HASH, pw):
            session.clear()
            session['admin'] = True
            return redirect(url_for('dashboard'))
        flash('비밀번호가 올바르지 않습니다.', 'error')
    return render_template('login.html')

@app.post('/teacher/logout')
@admin_required
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.get('/teacher')
@admin_required
def dashboard():
    rows,q = get_rows()
    return render_template('dashboard.html',rows=rows,q=q)

@app.get('/teacher/entry/<int:entry_id>')
@admin_required
def detail(entry_id):
    with connect() as con:
        entry = con.execute('SELECT * FROM submissions WHERE id=?',(entry_id,)).fetchone()
    if entry is None: abort(404)
    return render_template('detail.html',entry=entry)

@app.post('/teacher/entry/<int:entry_id>/delete')
@admin_required
def delete(entry_id):
    with connect() as con:
        con.execute('DELETE FROM submissions WHERE id=?',(entry_id,))
    flash('제출물이 삭제되었습니다.', 'success')
    return redirect(url_for('dashboard'))

@app.get('/teacher/export.xlsx')
@admin_required
def export_excel():
    rows,_ = get_rows()
    wb=Workbook(); ws=wb.active; ws.title='독후감'
    ws.append(['번호','학번','이름','책 제목','작가','전체 감상문','제출 일시'])
    for r in rows:
        # Prevent spreadsheet formula execution for untrusted student text.
        vals=[r['id'],r['student_id'],r['student_name'],r['book_title'],r['author'],r['review'],r['created_at']]
        ws.append(vals)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
    for cell in ws[1]: cell.fill=PatternFill('solid',fgColor='1D4ED8');cell.font=Font(color='FFFFFF',bold=True)
    for col,width in {'A':9,'B':18,'C':18,'D':35,'E':25,'F':85,'G':24}.items(): ws.column_dimensions[col].width=width
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    out=io.BytesIO(); wb.save(out);out.seek(0)
    return send_file(out,as_attachment=True,download_name='독후감_제출목록.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.get('/teacher/export.pdf')
@admin_required
def export_pdf():
    rows,_ = get_rows()
    font_candidates=[Path('/usr/share/fonts/truetype/nanum/NanumGothic.ttf'),Path('/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc')]
    font=next((p for p in font_candidates if p.exists()),None)
    if not font: abort(500,description='Korean PDF font missing')
    pdfmetrics.registerFont(TTFont('Korean',str(font)))
    styles=getSampleStyleSheet()
    title=ParagraphStyle('TitleKo',parent=styles['Title'],fontName='Korean',fontSize=17,leading=25,alignment=TA_CENTER)
    head=ParagraphStyle('HeadKo',fontName='Korean',fontSize=11,leading=18,spaceAfter=7)
    body=ParagraphStyle('BodyKo',fontName='Korean',fontSize=9,leading=15,spaceAfter=10,wordWrap='CJK')
    out=io.BytesIO();doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=40,leftMargin=40,topMargin=40,bottomMargin=40)
    story=[Paragraph('독후감 제출 목록',title),Spacer(1,16)]
    for r in rows:
        header=f"{escape(r['book_title'])} | {escape(r['student_id'])} {escape(r['student_name'])} | {escape(r['author'])} | {escape(r['created_at'])}"
        review=escape(r['review']).replace('\n','<br/>')
        story.extend([Paragraph(header,head),Paragraph(review,body),Spacer(1,12)])
    if not rows:story.append(Paragraph('제출된 독후감이 없습니다.',body))
    doc.build(story);out.seek(0)
    return send_file(out,as_attachment=True,download_name='독후감_제출목록.pdf',mimetype='application/pdf')

@app.get('/health')
def health():return {'status':'ok'}

if __name__=='__main__':app.run(debug=False)
